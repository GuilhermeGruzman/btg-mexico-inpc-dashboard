#!/usr/bin/env python3
"""
Mexico INPC dashboard updater

Workflow
--------
1) Query the official INEGI ca55/ca56 JSON services for the latest release.
2) Validate the 16-component release payload and cache new observations locally.
3) Combine the official historical seed with the JSON release cache.
4) Recalculate YoY, NSA, SA, 3M SAAR, 6M SAAR, seasonality and rankings.
5) Regenerate a standalone dashboard HTML.

The bundled XLSX files are historical seeds only. New releases do NOT require
downloading an INEGI spreadsheet. The final HTML embeds all dashboard data.

Final seasonal-adjustment choice: STL on log index levels for speed, reproducibility, and operational robustness.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import re
import shutil
import sys
import tempfile
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from statsmodels.tsa.seasonal import STL

MONTHLY_URL = "https://www.inegi.org.mx/app/tabulados/inp/default.aspx?nc=ca55_2018a&idrt=137&opc=t"
BIWEEKLY_URL = "https://www.inegi.org.mx/app/tabulados/inp/default.aspx?nc=ca56_2018a&idrt=137&opc=t"

MONTHLY_API_URL = "https://www.inegi.org.mx/app/tabulados/inp2/serviciocuadros/wsDataService.svc/obtienetabuladoinp/CA55_2018A/4/1"
BIWEEKLY_API_URL = "https://www.inegi.org.mx/app/tabulados/inp2/serviciocuadros/wsDataService.svc/obtienetabuladoinp/CA56_2018A/4/1"

COMPONENTS = [
    "CPI","Core","Goods","Food, Bvgs & Tobacco","Goods Ex-Food","Services",
    "Housing","Education","Other Services","Non Core","Agricultural",
    "Fruits & Vegetables","Meat & Eggs","Energy & Controlled Prices",
    "Energy","Controlled Prices"
]

# 2024 INPC basket weights used for decomposition display.
WEIGHTS = {
    "CPI":1.0,
    "Core":0.767415,
    "Goods":0.375338,
    "Food, Bvgs & Tobacco":0.172148,
    "Goods Ex-Food":0.203190,
    "Services":0.392077,
    "Housing":0.180550,
    "Education":0.025207,
    "Other Services":0.186321,
    "Non Core":0.232585,
    "Agricultural":0.106577,
    "Fruits & Vegetables":0.047789,
    "Meat & Eggs":0.058788,
    "Energy & Controlled Prices":0.126008,
    "Energy":0.080458,
    "Controlled Prices":0.045550,
}

MONTH_ES = {
    "Ene":1,"Feb":2,"Mar":3,"Abr":4,"May":5,"Jun":6,
    "Jul":7,"Ago":8,"Sep":9,"Oct":10,"Nov":11,"Dic":12
}
MONTH_EN = {1:"Jan",2:"Feb",3:"Mar",4:"Apr",5:"May",6:"Jun",7:"Jul",8:"Aug",9:"Sep",10:"Oct",11:"Nov",12:"Dec"}


def log(msg: str) -> None:
    print(f"[INPC] {msg}")


def fetch_xlsx_with_playwright(url: str, destination: Path) -> None:
    """Fetch the XLSX export from an INEGI tabulated-data page.

    INEGI does not always expose the export as a browser 'download' event.
    This implementation therefore listens to all network responses generated
    by the XLSX control and captures any response that is either:
      - XLSX content type;
      - an attachment whose filename ends in .xlsx/.xls; or
      - a response URL ending in .xlsx/.xls.

    It still supports a normal Playwright download event when INEGI uses one.
    """
    from playwright.sync_api import sync_playwright, TimeoutError as PlaywrightTimeoutError
    import time

    destination.parent.mkdir(parents=True, exist_ok=True)

    def looks_like_excel_response(response):
        try:
            headers = {k.lower(): v for k, v in response.headers.items()}
            ctype = headers.get("content-type", "").lower()
            dispo = headers.get("content-disposition", "").lower()
            rurl = response.url.lower().split("?")[0]
            return (
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in ctype
                or "application/vnd.ms-excel" in ctype
                or ".xlsx" in dispo
                or ".xls" in dispo
                or rurl.endswith(".xlsx")
                or rurl.endswith(".xls")
            )
        except Exception:
            return False

    with sync_playwright() as p:
        browser = None
        last_error = None
        for launch_kwargs in (
            {"channel": "chrome", "headless": True},
            {"headless": True},
        ):
            try:
                browser = p.chromium.launch(**launch_kwargs)
                break
            except Exception as exc:
                last_error = exc

        if browser is None:
            raise RuntimeError(
                "Could not launch Chrome/Chromium. Run: python -m playwright install chromium"
            ) from last_error

        try:
            context = browser.new_context(accept_downloads=True)
            page = context.new_page()

            captured = {"bytes": None, "url": None, "error": None}

            def on_response(response):
                if captured["bytes"] is not None:
                    return
                if not looks_like_excel_response(response):
                    return
                try:
                    body = response.body()
                    # XLSX files are ZIP containers and normally start with PK.
                    # Some servers may omit/alter headers, so also accept a
                    # sufficiently large body selected by XLSX metadata.
                    if body and (body[:2] == b"PK" or len(body) > 5000):
                        captured["bytes"] = body
                        captured["url"] = response.url
                except Exception as exc:
                    captured["error"] = str(exc)

            page.on("response", on_response)

            log(f"Opening INEGI page: {url}")
            page.goto(url, wait_until="domcontentloaded", timeout=45_000)
            page.wait_for_timeout(2000)
            log("INEGI page loaded. Looking for XLSX export control...")

            # Locate the Excel/XLSX control. INEGI may render this as an image,
            # anchor, button, input or element with an onclick handler.
            selectors = [
                "a:has(img[src*='ico_xlsx'])",
                "img[src*='ico_xlsx']",
                "a:has(img[alt*='Excel'])",
                "img[alt*='Excel']",
                "a:has(img[alt*='XLSX'])",
                "img[alt*='XLSX']",
                "[title*='Excel']",
                "[title*='XLSX']",
                "[onclick*='xlsx' i]",
                "[onclick*='excel' i]",
            ]

            target = None
            target_selector = None
            for selector in selectors:
                try:
                    loc = page.locator(selector)
                    if loc.count() > 0:
                        target = loc.first
                        target_selector = selector
                        break
                except Exception:
                    continue

            if target is None:
                # Diagnostic snapshot of likely export elements.
                diagnostic = page.locator("a,button,input,img").evaluate_all(
                    """els => els.map((e,i)=>({
                        i,
                        tag:e.tagName,
                        text:(e.innerText||e.value||e.alt||e.title||'').trim(),
                        href:e.href||'',
                        src:e.src||'',
                        onclick:e.getAttribute('onclick')||''
                    })).filter(x =>
                        /xls|excel|export/i.test(
                            [x.text,x.href,x.src,x.onclick].join(' ')
                        )
                    ).slice(0,30)"""
                )
                raise RuntimeError(
                    "Could not find INEGI XLSX export control. "
                    f"Candidate export elements: {diagnostic}"
                )

            # If the selector returned the icon itself, walk up to a clickable
            # ancestor. INEGI has changed this markup across versions.
            click_target = target
            try:
                clickable = target.locator(
                    "xpath=ancestor-or-self::*[self::a or self::button or @onclick][1]"
                )
                if clickable.count() > 0:
                    click_target = clickable.first
            except Exception:
                pass

            try:
                attrs = click_target.evaluate(
                    """e => ({
                        tag:e.tagName,
                        text:(e.innerText||e.value||e.alt||e.title||'').trim(),
                        href:e.href||'',
                        onclick:e.getAttribute('onclick')||''
                    })"""
                )
            except Exception:
                attrs = {"selector": target_selector}

            log(f"XLSX control found: {attrs}")
            log("Triggering XLSX export and monitoring INEGI network responses...")

            # Register the normal browser-download path as well. We do not block
            # on it, because INEGI may instead return the workbook through XHR/
            # iframe/form submission.
            normal_download = {"obj": None}
            page.on("download", lambda d: normal_download.__setitem__("obj", d))

            click_target.click(force=True, timeout=15_000)

            # Wait at most 35 seconds for either a normal download or a network
            # response containing the workbook.
            deadline = time.time() + 35
            last_log = 0
            while time.time() < deadline:
                if normal_download["obj"] is not None:
                    log("INEGI emitted a standard browser download.")
                    normal_download["obj"].save_as(str(destination))
                    return
                if captured["bytes"] is not None:
                    destination.write_bytes(captured["bytes"])
                    log(f"Captured XLSX from network response: {captured['url']}")
                    return

                page.wait_for_timeout(250)
                elapsed = int(35 - max(0, deadline - time.time()))
                if elapsed >= last_log + 5:
                    last_log = elapsed
                    log(f"Waiting for XLSX response... {elapsed}s")

            # Final diagnostic: collect URLs seen in resource performance entries.
            try:
                resources = page.evaluate(
                    """performance.getEntriesByType('resource')
                       .map(x=>x.name)
                       .filter(x=>/xls|excel|export|tabulad|inp/i.test(x))
                       .slice(-40)"""
                )
            except Exception:
                resources = []

            raise RuntimeError(
                "INEGI page was reached and the XLSX control was clicked, "
                "but no Excel response was captured within 35 seconds. "
                f"Control={attrs}. Relevant resource URLs={resources}. "
                f"Response read error={captured['error']}"
            )
        finally:
            browser.close()

def workbook_rows(path: Path):
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    return list(ws.iter_rows(values_only=True))


def validate_inegi_file(path: Path, frequency: str) -> str:
    rows = workbook_rows(path)
    if len(rows) < 20 or len(rows[4]) < 17:
        raise ValueError(f"Unexpected workbook structure: {path}")

    header_text = " ".join(str(x or "") for row in rows[:13] for x in row[:17])
    required_tokens = ["INEGI", "Índice Nacional de Precios al Consumidor", "Actualización de Canasta y Ponderadores 2024"]
    missing = [t for t in required_tokens if t not in header_text]
    if missing:
        raise ValueError(f"Unexpected INEGI metadata in {path}; missing: {missing}")
    if rows[7][0] != "Cifra" or any(str(rows[7][i]).strip() != "Índices" for i in range(1,17)):
        raise ValueError(f"Expected index-level INPC table, but metadata changed: {path}")

    latest = None
    if frequency == "monthly":
        pat = re.compile(r"(Ene|Feb|Mar|Abr|May|Jun|Jul|Ago|Sep|Oct|Nov|Dic)\s+(\d{4})$")
    else:
        pat = re.compile(r"([12])Q\s+(Ene|Feb|Mar|Abr|May|Jun|Jul|Ago|Sep|Oct|Nov|Dic)\s+(\d{4})$")
    for row in rows:
        if row and isinstance(row[0], str) and pat.fullmatch(row[0].strip()):
            latest = row[0].strip()
    if latest is None:
        raise ValueError(f"No dated INPC observations found in {path}")
    return latest


def file_sha256(path: Path) -> str:
    h=hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda:f.read(1024*1024),b""):
            h.update(chunk)
    return h.hexdigest()


def validate_parsed_frame(df: pd.DataFrame, frequency: str) -> list[str]:
    issues=[]
    if df.index.has_duplicates:
        issues.append("duplicate dates")
    if not df.index.is_monotonic_increasing:
        issues.append("dates are not strictly increasing")
    cols=[c for c in COMPONENTS if c in df.columns]
    if len(cols)!=16:
        issues.append(f"expected 16 components, found {len(cols)}")
    latest=df.iloc[-1]
    missing_latest=[c for c in cols if pd.isna(latest[c])]
    if missing_latest:
        issues.append("missing latest values: "+", ".join(missing_latest))
    for c in cols:
        vals=pd.to_numeric(df[c],errors="coerce").dropna()
        if (vals<=0).any():
            issues.append(f"non-positive index level in {c}")
            break
    if frequency=="biweekly" and "half" in df.columns:
        bad=set(pd.to_numeric(df["half"],errors="coerce").dropna().astype(int).unique())-{1,2}
        if bad:
            issues.append(f"invalid half identifiers: {sorted(bad)}")
    return issues


def revision_rows(old: pd.DataFrame, new: pd.DataFrame, tolerance: float=1e-10):
    common=old.index.intersection(new.index)
    out=[]
    for dt in common:
        for c in COMPONENTS:
            a=old.at[dt,c] if c in old.columns else np.nan
            b=new.at[dt,c] if c in new.columns else np.nan
            if pd.notna(a) and pd.notna(b) and abs(float(a)-float(b))>tolerance:
                out.append([dt.strftime("%Y-%m-%d"),c,float(a),float(b),float(b)-float(a)])
    return out


def write_revision_audit(path: Path, rows, label: str):
    path.parent.mkdir(parents=True,exist_ok=True)
    with path.open("a",newline="",encoding="utf-8") as f:
        w=csv.writer(f)
        if f.tell()==0:
            w.writerow(["detected_at","frequency","date","component","old_index","new_index","change"])
        ts=datetime.now().isoformat(timespec="seconds")
        for r in rows:
            w.writerow([ts,label]+r)


def download_latest(monthly_path: Path, biweekly_path: Path, root: Path) -> tuple[str, str]:
    with tempfile.TemporaryDirectory(prefix="inegi_inpc_") as td:
        td = Path(td)
        mtmp = td / "ca55_latest.xlsx"
        btmp = td / "ca56_latest.xlsx"

        log("Downloading official monthly table (ca55)...")
        fetch_xlsx_with_playwright(MONTHLY_URL, mtmp)
        mlast = validate_inegi_file(mtmp, "monthly")
        log("Downloading official biweekly table (ca56)...")
        fetch_xlsx_with_playwright(BIWEEKLY_URL, btmp)
        blast = validate_inegi_file(btmp, "biweekly")

        new_m=parse_monthly(mtmp); new_b=parse_biweekly(btmp)
        for label,df,freq in [("monthly",new_m,"monthly"),("biweekly",new_b,"biweekly")]:
            issues=validate_parsed_frame(df,freq)
            if issues:
                raise ValueError(f"Downloaded {label} data failed QA: {issues}")

        old_m=parse_monthly(monthly_path) if monthly_path.exists() else None
        old_b=parse_biweekly(biweekly_path) if biweekly_path.exists() else None
        if old_m is not None and new_m.index.max() < old_m.index.max():
            raise ValueError("Monthly download is older than the local validated history")
        if old_b is not None and new_b.index.max() < old_b.index.max():
            raise ValueError("Biweekly download is older than the local validated history")

        stamp=datetime.now().strftime("%Y%m%d_%H%M%S")
        archive=root/"archive"/stamp
        archive.mkdir(parents=True,exist_ok=True)
        if monthly_path.exists(): shutil.copy2(monthly_path,archive/monthly_path.name)
        if biweekly_path.exists(): shutil.copy2(biweekly_path,archive/biweekly_path.name)

        if old_m is not None:
            rev=revision_rows(old_m,new_m)
            if rev: write_revision_audit(root/"logs"/"revision_audit.csv",rev,"monthly")
        if old_b is not None:
            rev=revision_rows(old_b,new_b)
            if rev: write_revision_audit(root/"logs"/"revision_audit.csv",rev,"biweekly")

        shutil.copy2(mtmp, monthly_path)
        shutil.copy2(btmp, biweekly_path)
        return mlast, blast

def to_float(v):
    try:
        return float(v)
    except Exception:
        return np.nan


def parse_monthly(path: Path) -> pd.DataFrame:
    rows = workbook_rows(path)
    out=[]
    pat=re.compile(r"(Ene|Feb|Mar|Abr|May|Jun|Jul|Ago|Sep|Oct|Nov|Dic)\s+(\d{4})$")
    for row in rows:
        if not row or not isinstance(row[0],str):
            continue
        m=pat.fullmatch(row[0].strip())
        if not m:
            continue
        dt=pd.Timestamp(int(m.group(2)), MONTH_ES[m.group(1)], 15)
        vals=[to_float(row[i]) if i < len(row) else np.nan for i in range(1,17)]
        out.append([dt]+vals)
    df=pd.DataFrame(out,columns=["date"]+COMPONENTS).set_index("date").sort_index()
    if df.empty:
        raise ValueError("Monthly INEGI table parsed to zero rows")
    return df


def parse_biweekly(path: Path) -> pd.DataFrame:
    rows=workbook_rows(path)
    out=[]
    pat=re.compile(r"([12])Q\s+(Ene|Feb|Mar|Abr|May|Jun|Jul|Ago|Sep|Oct|Nov|Dic)\s+(\d{4})$")
    for row in rows:
        if not row or not isinstance(row[0],str):
            continue
        m=pat.fullmatch(row[0].strip())
        if not m:
            continue
        half=int(m.group(1)); mo=MONTH_ES[m.group(2)]; yr=int(m.group(3))
        dt=pd.Timestamp(yr,mo,15 if half==1 else 28)
        vals=[to_float(row[i]) if i < len(row) else np.nan for i in range(1,17)]
        out.append([dt,half]+vals)
    df=pd.DataFrame(out,columns=["date","half"]+COMPONENTS).set_index("date").sort_index()
    if df.empty:
        raise ValueError("Biweekly INEGI table parsed to zero rows")
    return df



def fetch_inegi_json(url: str) -> dict:
    """Fetch official INEGI tabulated-data JSON without launching a browser."""
    from playwright.sync_api import sync_playwright

    log(f"Querying INEGI JSON service: {url}")
    with sync_playwright() as p:
        request = p.request.new_context(
            extra_http_headers={
                "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/124 Safari/537.36",
                "Accept": "application/json, text/plain, */*",
                "Referer": "https://www.inegi.org.mx/",
                "Accept-Language": "es-MX,es;q=0.9,en;q=0.8",
            }
        )
        try:
            response = request.get(url, timeout=30_000)
            if not response.ok:
                raise RuntimeError(f"INEGI JSON returned HTTP {response.status}")
            data = response.json()
        finally:
            request.dispose()

    if not isinstance(data, dict):
        raise ValueError("Unexpected INEGI JSON payload type")
    return data


def parse_inegi_release(payload: dict, frequency: str) -> dict:
    """Validate and normalize the latest ca55/ca56 JSON release."""
    datos = payload.get("Datos")
    encab = payload.get("Encab")
    info = payload.get("InfoTab")
    available = payload.get("PeriodoDisponible")

    if not isinstance(datos, list) or len(datos) != 16:
        raise ValueError(f"Expected 16 INPC components, received {len(datos) if isinstance(datos,list) else 'invalid'}")
    if not isinstance(encab, list) or not encab:
        raise ValueError("INEGI JSON has no Encab block")
    if not isinstance(info, list) or not info:
        raise ValueError("INEGI JSON has no InfoTab block")

    title = str(info[0].get("titulo", ""))
    if "Índice Nacional de Precios al Consumidor" not in title and "Indice Nacional de Precios al Consumidor" not in title:
        raise ValueError(f"Unexpected INEGI table title: {title}")

    freq_expected = "Mensual" if frequency == "monthly" else "Quincenal"
    if isinstance(available, dict):
        freq_received = str(available.get("frecuencia", ""))
        if freq_received and freq_received != freq_expected:
            raise ValueError(f"Unexpected frequency: {freq_received}")

    period_text = str(encab[0].get("periodo_actual", "")).strip()
    if frequency == "monthly":
        m = re.fullmatch(r"(Ene|Feb|Mar|Abr|May|Jun|Jul|Ago|Sep|Oct|Nov|Dic)\s+(\d{4})", period_text)
        if not m:
            raise ValueError(f"Could not parse monthly period: {period_text}")
        dt = pd.Timestamp(int(m.group(2)), MONTH_ES[m.group(1)], 15)
        half = None
        cache_key = dt.strftime("%Y-%m")
    else:
        m = re.fullmatch(r"([12])Q\s+(Ene|Feb|Mar|Abr|May|Jun|Jul|Ago|Sep|Oct|Nov|Dic)\s+(\d{4})", period_text)
        if not m:
            raise ValueError(f"Could not parse biweekly period: {period_text}")
        half = int(m.group(1))
        dt = pd.Timestamp(int(m.group(3)), MONTH_ES[m.group(2)], 15 if half == 1 else 28)
        cache_key = f"{dt.strftime('%Y-%m')}-{half}"

    # INEGI occasionally publishes a duplicated/missing value in the metadata
    # field `orden` (e.g. ca56 can contain [..., 6, 6, 8, ...]) even though
    # all 16 component rows are present in the correct display sequence.
    # Therefore validate the structure without requiring an exact 1..16 vector.
    ordered = sorted(
        enumerate(datos),
        key=lambda pair: (int(pair[1].get("orden", 999)), pair[0])
    )
    ordered = [raw for _, raw in ordered]
    orders = [int(x.get("orden", -1)) for x in ordered]

    if len(orders) != 16:
        raise ValueError(f"Expected 16 component-order records, got {len(orders)}")
    if any(o < 1 or o > 16 for o in orders):
        raise ValueError(f"Invalid INEGI component-order metadata: {orders}")
    if any(orders[i] > orders[i+1] for i in range(len(orders)-1)):
        raise ValueError(f"Non-monotonic INEGI component ordering: {orders}")

    # Series IDs should still uniquely identify the 16 rows.
    series_ids = [str(x.get("serie", "")).strip() for x in ordered]
    if any(not s for s in series_ids) or len(set(series_ids)) != 16:
        raise ValueError(f"Invalid or duplicated INEGI series IDs: {series_ids}")

    rows = {}
    for component, raw in zip(COMPONENTS, ordered):
        try:
            period_change = float(raw["valor_mensual"])
            yoy = float(raw["valor_anual"])
            incidence = float(raw["valor_incidencia"]) if raw.get("valor_incidencia") not in (None, "") else None
        except Exception as exc:
            raise ValueError(f"Invalid numeric values for {component}: {raw}") from exc
        rows[component] = {
            "period": period_change,
            "yoy": yoy,
            "incidence": incidence,
            "series_id": str(raw.get("serie", "")),
            "description": str(raw.get("descripcion", "")),
        }

    return {
        "frequency": frequency,
        "period_text": period_text,
        "date": dt.strftime("%Y-%m-%d"),
        "half": half,
        "cache_key": cache_key,
        "rows": rows,
        "fetched_at": datetime.now().isoformat(timespec="seconds"),
        "source": "INEGI ca55 JSON service" if frequency == "monthly" else "INEGI ca56 JSON service",
    }


def load_release_cache(path: Path) -> dict:
    empty = {"monthly": {}, "biweekly": {}}
    if not path.exists():
        return empty
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
        if not isinstance(data, dict):
            return empty
        data.setdefault("monthly", {})
        data.setdefault("biweekly", {})
        return data
    except Exception:
        return empty


def save_release_cache(path: Path, cache: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(".tmp")
    tmp.write_text(json.dumps(cache, ensure_ascii=False, indent=2), encoding="utf-8")
    tmp.replace(path)


def _expected_next_date(prev: pd.Timestamp, frequency: str, next_half=None) -> pd.Timestamp:
    if frequency == "monthly":
        nxt = prev + pd.offsets.MonthBegin(1)
        return pd.Timestamp(nxt.year, nxt.month, 15)

    prev_half = 1 if prev.day <= 20 else 2
    if prev_half == 1:
        expected = pd.Timestamp(prev.year, prev.month, 28)
        if next_half is not None and next_half != 2:
            raise ValueError("Expected 2H after 1H in biweekly sequence")
        return expected

    nxt = prev + pd.offsets.MonthBegin(1)
    expected = pd.Timestamp(nxt.year, nxt.month, 15)
    if next_half is not None and next_half != 1:
        raise ValueError("Expected 1H after 2H in biweekly sequence")
    return expected


def apply_release_cache(df: pd.DataFrame, cache_records: dict, frequency: str) -> pd.DataFrame:
    """Append cached JSON releases to the historical index seed.

    The JSON service publishes period changes rounded to two decimals rather
    than index levels. For a genuinely new release we infer the next index
    level from the previous validated index and the official period change.
    """
    out = df.copy()

    records = []
    for _, rec in cache_records.items():
        try:
            dt = pd.Timestamp(rec["date"])
            records.append((dt, rec))
        except Exception:
            continue
    records.sort(key=lambda x: x[0])

    for dt, rec in records:
        half = rec.get("half")
        rows = rec.get("rows", {})
        if len(rows) != 16:
            continue

        if dt in out.index:
            # Historical seed is more precise than the rounded JSON variation.
            # Keep the official index and only use JSON as a consistency check.
            prev_candidates = out.index[out.index < dt]
            if len(prev_candidates):
                prev = prev_candidates.max()
                for c in COMPONENTS:
                    calc = (float(out.at[dt,c]) / float(out.at[prev,c]) - 1.0) * 100.0
                    official = float(rows[c]["period"])
                    if abs(calc - official) > 0.035:
                        raise ValueError(
                            f"JSON consistency check failed for {frequency} {rec.get('period_text')} {c}: "
                            f"seed={calc:.3f}% vs API={official:.3f}%"
                        )
            continue

        if dt < out.index.max():
            # Do not inject rounded derived levels into older historical gaps.
            continue

        prev = out.index.max()
        expected = _expected_next_date(prev, frequency, half)
        if dt != expected:
            raise ValueError(
                f"Cannot append {frequency} JSON release {rec.get('period_text')}: "
                f"expected next date {expected.date()}, got {dt.date()}"
            )

        vals = {}
        for c in COMPONENTS:
            rate = float(rows[c]["period"]) / 100.0
            vals[c] = float(out.at[prev,c]) * (1.0 + rate)

        if frequency == "monthly":
            new_row = pd.DataFrame([vals], index=[dt])
        else:
            new_row = pd.DataFrame([{"half": int(half), **vals}], index=[dt])

        out = pd.concat([out, new_row]).sort_index()

        # Validate implied YoY against the official published YoY. The API
        # values are rounded, so use a small tolerance.
        lag = 12 if frequency == "monthly" else 24
        if len(out) > lag:
            pos = out.index.get_loc(dt)
            if isinstance(pos, (int, np.integer)) and pos >= lag:
                lag_dt = out.index[pos-lag]
                for c in COMPONENTS:
                    implied = (float(out.at[dt,c]) / float(out.at[lag_dt,c]) - 1.0) * 100.0
                    official_yoy = float(rows[c]["yoy"])
                    if abs(implied - official_yoy) > 0.08:
                        raise ValueError(
                            f"YoY validation failed for {frequency} {rec.get('period_text')} {c}: "
                            f"implied={implied:.3f}% vs API={official_yoy:.3f}%"
                        )

    return out


def update_release_cache(root: Path) -> tuple[str, str]:
    """Fetch latest ca55/ca56 JSON releases and persist them in a local cache."""
    cache_path = root / "data" / "inegi_release_cache.json"
    cache = load_release_cache(cache_path)

    log("Fetching official monthly release from INEGI JSON service...")
    monthly_release = parse_inegi_release(fetch_inegi_json(MONTHLY_API_URL), "monthly")
    log(f"Monthly API release: {monthly_release['period_text']}")

    log("Fetching official biweekly release from INEGI JSON service...")
    biweekly_release = parse_inegi_release(fetch_inegi_json(BIWEEKLY_API_URL), "biweekly")
    log(f"Biweekly API release: {biweekly_release['period_text']}")

    cache["monthly"][monthly_release["cache_key"]] = monthly_release
    cache["biweekly"][biweekly_release["cache_key"]] = biweekly_release
    save_release_cache(cache_path, cache)

    return monthly_release["period_text"], biweekly_release["period_text"]



def pct(s: pd.Series, periods: int=1) -> pd.Series:
    return s.pct_change(periods=periods, fill_method=None)*100.0


def stl_sa(series: pd.Series, period: int) -> pd.Series:
    s=series.dropna().astype(float)
    if len(s)<period*3:
        return pd.Series(index=series.index,dtype=float)
    fit=STL(np.log(s.values),period=period,robust=True).fit()
    sa=np.exp(fit.trend+fit.resid)
    return pd.Series(sa,index=s.index).reindex(series.index)


def iso_dates(index):
    return [d.strftime("%Y-%m-%d") for d in index]


def js_series(name,s):
    return {"name":name,"x":iso_dates(s.index),"y":[None if pd.isna(v) else float(v) for v in s.tolist()]}


def js_chart(title,ctype,series):
    return {"title":title,"type":ctype,"series":series}


def normalized_contrib(yoy_df,children,parent):
    raw=pd.DataFrame(index=yoy_df.index)
    pw=WEIGHTS[parent]
    for c in children:
        raw[c]=(WEIGHTS[c]/pw)*yoy_df[c]
    factor=yoy_df[parent]/raw.sum(axis=1).replace(0,np.nan)
    return raw.mul(factor,axis=0)



def _period_key(value) -> str:
    if pd.isna(value):
        return ""
    if isinstance(value, pd.Timestamp):
        return value.strftime("%Y-%m")
    s=str(value).strip()
    m=re.search(r"(\d{4})[-/](\d{1,2})",s)
    if m:
        return f"{int(m.group(1)):04d}-{int(m.group(2)):02d}"
    return s[:7]


def load_btg_projections(path: Path):
    """Read pre-release desk forecasts and return (payload, warnings)."""
    empty={"Monthly":{},"1H":{},"2H":{}}
    warnings=[]
    if not path.exists():
        warnings.append(f"BTG projection file not found: {path.name}")
        return empty,warnings
    try:
        df=pd.read_excel(path,sheet_name="Projections")
    except Exception as exc:
        warnings.append(f"Could not read BTG projection file: {exc}")
        return empty,warnings

    required={"Release","Reference Period","Component","Period Projection (%)"}
    missing=required-set(df.columns)
    if missing:
        warnings.append(f"Projection file missing columns: {sorted(missing)}")
        return empty,warnings

    out={"Monthly":{},"1H":{},"2H":{}}
    seen=set()
    for i,row in df.iterrows():
        rel=str(row.get("Release","")).strip()
        period=_period_key(row.get("Reference Period"))
        component=str(row.get("Component","")).strip()
        if not rel and not period and not component:
            continue
        if rel not in out:
            warnings.append(f"Row {i+2}: invalid Release '{rel}'")
            continue
        if not period or component not in COMPONENTS:
            warnings.append(f"Row {i+2}: invalid period/component")
            continue
        key=(rel,period,component)
        if key in seen:
            warnings.append(f"Duplicate projection: {rel} {period} {component}; last row wins")
        seen.add(key)

        p=row.get("Period Projection (%)")
        if pd.isna(p):
            pval=None
        else:
            try: pval=float(p)
            except Exception:
                warnings.append(f"Row {i+2}: invalid numeric projection")
                pval=None
        if pval is not None and abs(pval)>10:
            warnings.append(f"Row {i+2}: unusually large period projection ({pval:.2f}%)")

        y=row.get("YoY Projection (%)") if "YoY Projection (%)" in df.columns else None
        yval=None if y is None or pd.isna(y) else float(y)
        notes=str(row.get("Notes","") or "")
        timestamp=row.get("Forecast Timestamp") if "Forecast Timestamp" in df.columns else None
        owner=row.get("Analyst / Source") if "Analyst / Source" in df.columns else None

        rec=out[rel].setdefault(period,{"demo":True,"rows":{},"forecast_timestamp":None,"analyst_source":None})
        rec["rows"][component]={"period":pval,"yoy":yval}
        if timestamp is not None and not pd.isna(timestamp): rec["forecast_timestamp"]=str(timestamp)
        if owner is not None and not pd.isna(owner): rec["analyst_source"]=str(owner)
        if "ILLUSTRATIVE" not in notes.upper() and "DEMO" not in notes.upper():
            rec["demo"]=False

    for rel,pmap in out.items():
        for period,rec in pmap.items():
            present=[c for c,v in rec["rows"].items() if v.get("period") is not None]
            if len(present)<16:
                warnings.append(f"{rel} {period}: {len(present)}/16 period projections populated")
            if "CPI" not in present or "Core" not in present:
                warnings.append(f"{rel} {period}: CPI/Core projection missing")
    return out,warnings



def build_payloads(monthly_path: Path, biweekly_path: Path, projections=None, release_cache=None):
    midx=parse_monthly(monthly_path)
    bidx=parse_biweekly(biweekly_path)
    release_cache=release_cache or {"monthly":{},"biweekly":{}}
    midx=apply_release_cache(midx,release_cache.get("monthly",{}),"monthly")
    bidx=apply_release_cache(bidx,release_cache.get("biweekly",{}),"biweekly")

    myoy=midx[COMPONENTS].apply(lambda s:pct(s,12))
    mmom=midx[COMPONENTS].apply(lambda s:pct(s,1))
    msa=pd.DataFrame(index=midx.index)
    msa1=pd.DataFrame(index=midx.index)
    m3=pd.DataFrame(index=midx.index)
    m6=pd.DataFrame(index=midx.index)
    for c in COMPONENTS:
        sa=stl_sa(midx[c],12)
        msa[c]=sa; msa1[c]=pct(sa,1)
        m3[c]=((sa/sa.shift(3))**4-1)*100
        m6[c]=((sa/sa.shift(6))**2-1)*100

    bvals=bidx[COMPONENTS]
    bnsa=bvals.apply(lambda s:pct(s,1))
    byoy=bvals.apply(lambda s:pct(s,24))
    bsa1=pd.DataFrame(index=bidx.index)
    for c in COMPONENTS:
        bsa1[c]=pct(stl_sa(bidx[c],24),1)

    half={}
    for h in [1,2]:
        idxh=bidx[bidx["half"]==h][COMPONENTS]
        hyoy=idxh.apply(lambda s:pct(s,12))
        h3=pd.DataFrame(index=idxh.index); h6=pd.DataFrame(index=idxh.index)
        for c in COMPONENTS:
            sa=stl_sa(idxh[c],12)
            h3[c]=((sa/sa.shift(3))**4-1)*100
            h6[c]=((sa/sa.shift(6))**2-1)*100
        half[h]={"idx":idxh,"yoy":hyoy,"m3":h3,"m6":h6}

    headline=["Goods","Services","Agricultural","Energy & Controlled Prices"]
    core=["Food, Bvgs & Tobacco","Goods Ex-Food","Housing","Education","Other Services"]
    noncore=["Agricultural","Energy","Controlled Prices"]
    mh=normalized_contrib(myoy,headline,"CPI")
    mc=normalized_contrib(myoy,core,"Core")
    mn=normalized_contrib(myoy,noncore,"Non Core")

    projections=projections or {"Monthly":{},"1H":{},"2H":{}}

    def official_release_rows(release,date,half_value=None):
        if release=="Monthly":
            rec=release_cache.get("monthly",{}).get(date.strftime("%Y-%m"),{})
        else:
            h=int(half_value if half_value is not None else (1 if release=="1H" else 2))
            rec=release_cache.get("biweekly",{}).get(f"{date.strftime('%Y-%m')}-{h}",{})
        return rec.get("rows",{}) if isinstance(rec,dict) else {}

    def projection_for(release,date,component):
        key=date.strftime("%Y-%m")
        rec=projections.get(release,{}).get(key,{})
        return rec.get("rows",{}).get(component,{})

    def aligned_projection(release,date,component,actual_period,actual_yoy):
        pr=projection_for(release,date,component)
        pp=pr.get("period")
        if pp is None or pd.isna(actual_period) or pd.isna(actual_yoy):
            return pp, pr.get("yoy")
        # Keep YoY mechanically consistent with the projected period change.
        # (1 + projected YoY) = (1 + projected period rate) * I_{t-1}/I_{t-12}
        base_ratio=(1.0+float(actual_yoy)/100.0)/(1.0+float(actual_period)/100.0)
        pyoy=((1.0+float(pp)/100.0)*base_ratio-1.0)*100.0
        return float(pp), float(pyoy)

    def monthly_table():
        dates=midx.index[-3:]
        rows=[]
        latest=dates[-1]
        incidence_by_date={d:official_release_rows("Monthly",d) for d in dates}
        for c in COMPONENTS:
            pp,pyoy=aligned_projection(
                "Monthly",latest,c,
                mmom.loc[latest,c],
                myoy.loc[latest,c]
            )
            contrib=[]
            for d in dates:
                inc=incidence_by_date[d].get(c,{}).get("incidence")
                contrib.append(None if inc is None else float(inc)/100.0)
            rows.append({
                "component":c,"weight":WEIGHTS[c],
                "p":[None if pd.isna(v) else float(v) for v in (mmom.loc[dates,c]/100)],
                "cons":None if pp is None else float(pp)/100.0,
                "contrib":contrib,
                "yoy":[None if pd.isna(v) else float(v) for v in (myoy.loc[dates,c]/100)],
                "ycons":None if pyoy is None else float(pyoy)/100.0
            })
        return {"dates":[f"{MONTH_EN[d.month]} {d.year}" for d in dates],"rows":rows}

    def blabel(d,h): return f"{'1H' if h==1 else '2H'} {MONTH_EN[d.month]} {d.year}"
    def biweekly_table():
        dates=bidx.index[-3:]
        rows=[]
        latest=dates[-1]
        latest_half=int(bidx.loc[latest,"half"])
        release="1H" if latest_half==1 else "2H"
        incidence_by_date={
            d:official_release_rows(
                "1H" if int(bidx.loc[d,"half"])==1 else "2H",
                d,
                int(bidx.loc[d,"half"])
            )
            for d in dates
        }
        for c in COMPONENTS:
            pp,pyoy=aligned_projection(
                release,latest,c,
                bnsa.loc[latest,c],
                byoy.loc[latest,c]
            )
            contrib=[]
            for d in dates:
                inc=incidence_by_date[d].get(c,{}).get("incidence")
                contrib.append(None if inc is None else float(inc)/100.0)
            rows.append({
                "component":c,"weight":WEIGHTS[c],
                "p":[None if pd.isna(v) else float(v) for v in (bnsa.loc[dates,c]/100)],
                "cons":None if pp is None else float(pp)/100.0,
                "contrib":contrib,
                "yoy":[None if pd.isna(v) else float(v) for v in (byoy.loc[dates,c]/100)],
                "ycons":None if pyoy is None else float(pyoy)/100.0
            })
        return {"dates":[blabel(d,int(bidx.loc[d,"half"])) for d in dates],"rows":rows}

    monthly_charts=[
        js_chart("México - CPI (YoY)","line",[js_series("Headline",myoy["CPI"]),js_series("Core",myoy["Core"])]),
        js_chart("Mexico - CPI Headline","bar",[js_series("3M SAAR",m3["CPI"]),js_series("6M SAAR",m6["CPI"])]),
        js_chart("Mexico - CPI Core","bar",[js_series("3M SAAR",m3["Core"]),js_series("6M SAAR",m6["Core"])]),
        js_chart("Mexico - Services","bar",[js_series("3M SAAR",m3["Services"]),js_series("6M SAAR",m6["Services"])]),
        js_chart("Mexico - Goods","line",[js_series("3M SAAR",m3["Goods"]),js_series("6M SAAR",m6["Goods"])]),
        js_chart("Mexico - CPI Goods (YoY)","line",[js_series("Goods",myoy["Goods"]),js_series("Food, Bvgs & Tobacco",myoy["Food, Bvgs & Tobacco"]),js_series("Goods Ex-Food",myoy["Goods Ex-Food"])]),
        js_chart("Mexico - CPI Services (YoY)","line",[js_series("Other Services",myoy["Other Services"]),js_series("CPI Services",myoy["Services"])]),
        js_chart("Mexico - Other Services","bar",[js_series("3M SAAR",m3["Other Services"]),js_series("6M SAAR",m6["Other Services"])]),
        js_chart("Mexico - Energy & Controlled Prices","bar",[js_series("3M SAAR",m3["Energy & Controlled Prices"]),js_series("6M SAAR",m6["Energy & Controlled Prices"])]),
        js_chart("Mexico - CPI Headline (YoY Breakdown)","bar",[*[js_series(c,mh[c]) for c in headline],js_series("CPI",myoy["CPI"])]),
        js_chart("Mexico - CPI Core (YoY Breakdown)","bar",[*[js_series(c,mc[c]) for c in core],js_series("Core",myoy["Core"])]),
        js_chart("Mexico - CPI (YoY)","line",[js_series("Goods",myoy["Goods"]),js_series("Services",myoy["Services"])]),
        js_chart("Mexico: Non Core Breakdown YoY","bar",[*[js_series(c,mn[c]) for c in noncore],js_series("Non Core",myoy["Non Core"])])
    ]

    def half_charts(h):
        label="Mid-Month" if h==1 else "End Month"
        d=half[h]
        hc=normalized_contrib(d["yoy"],headline,"CPI")
        cc=normalized_contrib(d["yoy"],core,"Core")
        return [
            js_chart(f"Mexico - {label} CPI (YoY)","line",[js_series("Headline",d["yoy"]["CPI"]),js_series("Core",d["yoy"]["Core"])]),
            js_chart(f"Mexico - {label} CPI Core Goods & Services (YoY)","line",[js_series("Goods",d["yoy"]["Goods"]),js_series("Services",d["yoy"]["Services"])]),
            js_chart(f"Mexico - {label} CPI Headline","bar",[js_series("3M SAAR",d["m3"]["CPI"]),js_series("6M SAAR",d["m6"]["CPI"])]),
            js_chart(f"Mexico - {label} CPI Core","bar",[js_series("3M SAAR",d["m3"]["Core"]),js_series("6M SAAR",d["m6"]["Core"])]),
            js_chart(f"Mexico - {label} Goods","bar",[js_series("3M SAAR",d["m3"]["Goods"]),js_series("6M SAAR",d["m6"]["Goods"])]),
            js_chart(f"Mexico - {label} Services","bar",[js_series("3M SAAR",d["m3"]["Services"]),js_series("6M SAAR",d["m6"]["Services"])]),
            js_chart(f"Mexico - {label} Other Services","bar",[js_series("3M SAAR",d["m3"]["Other Services"]),js_series("6M SAAR",d["m6"]["Other Services"])]),
            js_chart(f"Mexico - {label} Food, Bvgs & Tobacco","bar",[js_series("3M SAAR",d["m3"]["Food, Bvgs & Tobacco"]),js_series("6M SAAR",d["m6"]["Food, Bvgs & Tobacco"])]),
            js_chart(f"Mexico - {label} CPI Headline (YoY Breakdown)","bar",[*[js_series(c,hc[c]) for c in headline],js_series("CPI",d["yoy"]["CPI"])]),
            js_chart(f"Mexico - {label} CPI Core (YoY Breakdown)","bar",[*[js_series(c,cc[c]) for c in core],js_series("Core",d["yoy"]["Core"])])
        ]

    D={"monthly":{"table":monthly_table(),"charts":monthly_charts},"biweekly":{"table":biweekly_table(),"charts":half_charts(1)+half_charts(2)}}

    season={"Monthly":{c:[] for c in COMPONENTS},"1H":{c:[] for c in COMPONENTS},"2H":{c:[] for c in COMPONENTS}}
    for c in COMPONENTS:
        for dt,v in mmom[c].items():
            if pd.notna(v): season["Monthly"][c].append({"year":int(dt.year),"month":int(dt.month),"value":round(float(v),6)})
        for dt,v in bnsa[c].items():
            if pd.notna(v):
                rel="1H" if int(bidx.loc[dt,"half"])==1 else "2H"
                season[rel][c].append({"year":int(dt.year),"month":int(dt.month),"value":round(float(v),6)})

    ranking={"data":{"Monthly":{},"1H":{},"2H":{}},"hist":{"Monthly":{},"1H":{},"2H":{}}}
    for dt in midx.index:
        key=dt.strftime("%Y-%m"); ranking["data"]["Monthly"][key]={}
        for c in COMPONENTS:
            vals={"yoy":myoy.at[dt,c],"nsa":mmom.at[dt,c],"sa":msa1.at[dt,c],"3m":m3.at[dt,c],"6m":m6.at[dt,c]}
            ranking["data"]["Monthly"][key][c]={k:(None if pd.isna(v) else float(v)) for k,v in vals.items()}
    for h,rel in [(1,"1H"),(2,"2H")]:
        idxh=half[h]["idx"]
        for dt in idxh.index:
            key=dt.strftime("%Y-%m"); ranking["data"][rel][key]={}
            for c in COMPONENTS:
                vals={"yoy":half[h]["yoy"].at[dt,c],"nsa":bnsa.at[dt,c],"sa":bsa1.at[dt,c],"3m":half[h]["m3"].at[dt,c],"6m":half[h]["m6"].at[dt,c]}
                ranking["data"][rel][key][c]={k:(None if pd.isna(v) else float(v)) for k,v in vals.items()}
    for rel in ["Monthly","1H","2H"]:
        for c in COMPONENTS: ranking["hist"][rel][c]={m:[] for m in ["yoy","nsa","sa","3m","6m"]}
        for _, cmap in ranking["data"][rel].items():
            for c,vals in cmap.items():
                for m,v in vals.items():
                    if v is not None and math.isfinite(v): ranking["hist"][rel][c][m].append(v)

    return D,season,ranking,midx,bidx


def replace_js_object(html: str, marker: str, obj) -> str:
    p=html.find(marker)
    if p<0: raise RuntimeError(f"Template is missing {marker}")
    start=p+len(marker)
    _, consumed=json.JSONDecoder().raw_decode(html[start:])
    end=start+consumed
    return html[:start]+json.dumps(obj,separators=(",",":"),ensure_ascii=False)+html[end:]



def enforce_dashboard_ui(html: str) -> str:
    # 1) Global demo note at top
    if ".global-demo-note{" not in html:
        css = (
            ".global-demo-note{display:none;font-size:10px;color:#765000;background:#FFF7DF;"
            "border:1px solid #E5C76B;border-left:4px solid #D5A100;border-radius:4px;"
            "padding:8px 10px;margin:0 0 10px 0;line-height:1.45;}"
            ".contrib-above-btg{color:#B42318!important;font-weight:700!important;}"
        )
        html = html.replace("</style>", css + "</style>", 1)

    if 'id="global-demo-note"' not in html:
        p = html.find('<div class="tabs">')
        if p >= 0:
            html = html[:p] + '<div id="global-demo-note" class="global-demo-note"></div>' + html[p:]

    # Remove in-panel demo banner
    html = re.sub(
        r"const demoBanner=projection\.demo[\s\S]*?:\s*'';",
        "const demoBanner='';",
        html,
        count=1
    )

    if "function renderGlobalDemoNote()" not in html:
        fn = (
            "\nfunction renderGlobalDemoNote(){\n"
            " const el=document.getElementById('global-demo-note');\n"
            " if(!el) return;\n"
            " let hasDemo=false;\n"
            " for(const rel of Object.keys(BTG_PROJECTIONS||{})){\n"
            "  for(const rec of Object.values(BTG_PROJECTIONS[rel]||{})){\n"
            "   if(rec && rec.demo){hasDemo=true;break;}\n"
            "  }\n"
            "  if(hasDemo) break;\n"
            " }\n"
            " if(hasDemo){\n"
            "  el.style.display='block';\n"
            "  el.innerHTML='<b>Illustrative BTG projections — demo only.</b> Current forecast figures are placeholders used to demonstrate the release-surprise workflow. Replace them in <b>btg_projections.xlsx</b> with the desk’s pre-release estimates before operational use.';\n"
            " }else{el.style.display='none';el.innerHTML='';}\n"
            "}\n"
        )
        p = html.find("function renderMonthlyPanel(){")
        if p >= 0:
            html = html[:p] + fn + html[p:]

    if "renderGlobalDemoNote();" not in html:
        init = "try{ renderMonthlyPanel(); }catch(err){ console.error('Monthly panel render failed:',err); }"
        if init in html:
            html = html.replace(
                init,
                "try{ renderGlobalDemoNote(); }catch(err){ console.error('Demo note render failed:',err); }\n" + init,
                1
            )

    # 2) Table: Dif. BTG cell is red when the contribution surprise is positive
    html = html.replace(
        "<td class=\"${diffContrib!==null&&diffContrib>0?'contrib-above-btg':''}\">${cellNum(actualContrib)}</td><td>${cellNum(diffContrib)}</td>",
        "<td>${cellNum(actualContrib)}</td><td class=\"${diffContrib!==null&&diffContrib>0?'contrib-above-btg':''}\">${cellNum(diffContrib)}</td>"
    )
    html = html.replace(
        "<td>${cellNum(actualContrib)}</td><td>${cellNum(diffContrib)}</td>",
        "<td>${cellNum(actualContrib)}</td><td class=\"${diffContrib!==null&&diffContrib>0?'contrib-above-btg':''}\">${cellNum(diffContrib)}</td>"
    )

    # 3) Ranking colors / bold labels only
    html = re.sub(
        r"const rankingColorMap=\{[\s\S]*?\};\s*const barColors=plotRows\.map\(r=>rankingColorMap\[r\.component\]\s*\|\|\s*'#[A-Fa-f0-9]+'\);",
        "const barColors=plotRows.map(r=>r.value<0?'#C94C4C':'#2F80C8');\n"
        " const keyComponents=new Set(['CPI','Core','Goods','Services']);\n"
        " const rankDisplayName=c=>c==='CPI'?'Headline':c;\n"
        " const rankTickText=plotRows.map(r=>{const label=rankDisplayName(r.component);return keyComponents.has(r.component)?`<b>${label}</b>`:label;});",
        html,
        count=1
    )

    if "const keyComponents=new Set(['CPI','Core','Goods','Services']);" not in html:
        html = re.sub(
            r"const barColors=plotRows\.map\(r=>[^;]+;",
            "const barColors=plotRows.map(r=>r.value<0?'#C94C4C':'#2F80C8');\n"
            " const keyComponents=new Set(['CPI','Core','Goods','Services']);\n"
            " const rankDisplayName=c=>c==='CPI'?'Headline':c;\n"
            " const rankTickText=plotRows.map(r=>{const label=rankDisplayName(r.component);return keyComponents.has(r.component)?`<b>${label}</b>`:label;});",
            html,
            count=1
        )

    html = html.replace("y:plotRows.map(r=>r.component),", "y:plotRows.map((r,i)=>String(i)),")

    old_hover = (
        "customdata:plotRows.map(r=>[\n"
        "       r.previous,\n"
        "       r.delta\n"
        "     ]),\n"
        "     hovertemplate:\n"
        "       '<b>%{y}</b><br>'+"
    )
    new_hover = (
        "customdata:plotRows.map(r=>[\n"
        "       r.previous,\n"
        "       r.delta,\n"
        "       rankDisplayName(r.component)\n"
        "     ]),\n"
        "     hovertemplate:\n"
        "       '<b>%{customdata[2]}</b><br>'+"
    )
    html = html.replace(old_hover, new_hover)

    rank_pos = html.find("function renderRanking(){")
    if rank_pos >= 0 and "ticktext:rankTickText" not in html[rank_pos:rank_pos+12000]:
        yaxis_pos = html.find("yaxis:{", rank_pos)
        if yaxis_pos >= 0:
            end = html.find("},", yaxis_pos)
            if end >= 0:
                html = html[:end] + ",tickmode:'array',tickvals:plotRows.map((r,i)=>String(i)),ticktext:rankTickText" + html[end:]

    # Remove Plotly trace-name box ("trace 0") from Ranking hover.
    rank_pos = html.find("function renderRanking(){")
    if rank_pos >= 0:
        rank_end = html.find("function ", rank_pos + 30)
        if rank_end < 0:
            rank_end = len(html)
        rank_block = html[rank_pos:rank_end]
        # Add <extra></extra> to the last line of the hovertemplate if not already present.
        if "<extra></extra>" not in rank_block and "hovertemplate:" in rank_block:
            rank_block = rank_block.replace(
                "'Δ: %{customdata[1]:+.2f}%<br>'",
                "'Δ: %{customdata[1]:+.2f}%<extra></extra>'"
            )
            rank_block = rank_block.replace(
                "'Δ: %{customdata[1]:+.2f}%'",
                "'Δ: %{customdata[1]:+.2f}%<extra></extra>'"
            )
            html = html[:rank_pos] + rank_block + html[rank_end:]

    # 4) Surprise decomposition with official realized incidence when available
    old = (
        "const rows=Object.keys(projection.rows).map(component=>{\n"
        "     const a=actual?.[component]?.nsa;\n"
        "     const ay=actual?.[component]?.yoy;\n"
        "     const p=projection.rows?.[component]?.period;\n"
        "     const py=projection.rows?.[component]?.yoy;\n"
        "     const surprise=Number.isFinite(a)&&Number.isFinite(p)?a-p:null;\n"
        "     const yoySurprise=Number.isFinite(ay)&&Number.isFinite(py)?ay-py:null;\n"
        "     const weight=SURPRISE_WEIGHTS[component];\n"
        "     const contrib=Number.isFinite(surprise)&&Number.isFinite(weight)?surprise*weight*100:null;"
    )
    if old in html:
        new = (
            "const officialRelease=INEGI_RELEASES?.[release]?.[date] || null;\n"
            "   const rows=Object.keys(projection.rows).map(component=>{\n"
            "     const a=actual?.[component]?.nsa;\n"
            "     const ay=actual?.[component]?.yoy;\n"
            "     const p=projection.rows?.[component]?.period;\n"
            "     const py=projection.rows?.[component]?.yoy;\n"
            "     const surprise=Number.isFinite(a)&&Number.isFinite(p)?a-p:null;\n"
            "     const yoySurprise=Number.isFinite(ay)&&Number.isFinite(py)?ay-py:null;\n"
            "     const weight=SURPRISE_WEIGHTS[component];\n"
            "     const officialInc=officialRelease?.rows?.[component]?.incidence;\n"
            "     const realizedContribution=Number.isFinite(officialInc)?officialInc:(Number.isFinite(a)&&Number.isFinite(weight)?a*weight:null);\n"
            "     const projectedContribution=Number.isFinite(p)&&Number.isFinite(weight)?p*weight:null;\n"
            "     const contrib=Number.isFinite(realizedContribution)&&Number.isFinite(projectedContribution)?(realizedContribution-projectedContribution)*100:null;"
        )
        html = html.replace(old, new, 1)
        html = html.replace(
            "return {component,a,ay,p,py,surprise,yoySurprise,contrib};",
            "return {component,a,ay,p,py,surprise,yoySurprise,contrib,realizedContribution,projectedContribution};",
            1
        )

    html = html.replace(
        "(Realized − BTG) × weight · contribution in basis points",
        "INEGI realized incidence − BTG projected contribution · basis points"
    )

    return html


def rebuild_dashboard(template: Path, output: Path, monthly_path: Path, biweekly_path: Path, projections_path: Path, root: Path):
    projections,projection_warnings=load_btg_projections(projections_path)
    release_cache=load_release_cache(root/"data"/"inegi_release_cache.json")
    D,season,ranking,midx,bidx=build_payloads(monthly_path,biweekly_path,projections,release_cache)

    # Align projection YoY with the period projection and realized historical index path.
    for rel in ["Monthly","1H","2H"]:
        for period,rec in projections.get(rel,{}).items():
            if period not in ranking["data"].get(rel,{}):
                continue
            for comp,row in rec.get("rows",{}).items():
                actual=ranking["data"][rel][period].get(comp,{})
                pp=row.get("period")
                ap=actual.get("nsa")
                ay=actual.get("yoy")
                if pp is None or ap is None or ay is None:
                    continue
                base_ratio=(1.0+float(ay)/100.0)/(1.0+float(ap)/100.0)
                row["yoy"]=((1.0+float(pp)/100.0)*base_ratio-1.0)*100.0

    official_payload={"Monthly":{},"1H":{},"2H":{}}
    for key,rec in release_cache.get("monthly",{}).items():
        period=str(key)[:7]
        official_payload["Monthly"][period]=rec
    for key,rec in release_cache.get("biweekly",{}).items():
        half=int(rec.get("half") or (1 if str(key).endswith("-1") else 2))
        period=str(key)[:7]
        official_payload["1H" if half==1 else "2H"][period]=rec

    html=template.read_text(encoding="utf-8")
    html=enforce_dashboard_ui(html)

    # Backward compatibility: older dashboard templates do not yet contain
    # the INEGI_RELEASES payload. Inject the placeholder automatically so
    # users can replace only update_dashboard.py without also replacing the template.
    if "const INEGI_RELEASES=" not in html:
        anchor="const BTG_PROJECTIONS="
        p=html.find(anchor)
        if p<0:
            raise RuntimeError("Template is missing const BTG_PROJECTIONS=; cannot inject INEGI_RELEASES")
        start=p+len(anchor)
        _,consumed=json.JSONDecoder().raw_decode(html[start:])
        end=start+consumed
        semi=end
        while semi < len(html) and html[semi].isspace():
            semi += 1
        if semi < len(html) and html[semi]==";":
            semi += 1
        injection='\nconst INEGI_RELEASES={"Monthly":{},"1H":{},"2H":{}};'
        html=html[:semi]+injection+html[semi:]

    html=replace_js_object(html,"const D=",D)
    html=replace_js_object(html,"const SEASONALITY_DATA=",season)
    html=replace_js_object(html,"const RANKING_DATA=",ranking)
    html=replace_js_object(html,"const BTG_PROJECTIONS=",projections)
    html=replace_js_object(html,"const INEGI_RELEASES=",official_payload)
    output.parent.mkdir(parents=True,exist_ok=True)
    output.write_text(html,encoding="utf-8")
    return midx,bidx,projection_warnings


def write_qa_report(root: Path, monthly_path: Path, biweekly_path: Path, midx: pd.DataFrame, bidx: pd.DataFrame, projection_warnings):
    issues=[]
    issues += ["Monthly: "+x for x in validate_parsed_frame(midx,"monthly")]
    issues += ["Biweekly: "+x for x in validate_parsed_frame(bidx,"biweekly")]
    weight_checks={
        "Core + Non Core": WEIGHTS["Core"]+WEIGHTS["Non Core"],
        "Goods + Services": WEIGHTS["Goods"]+WEIGHTS["Services"],
        "Core children": sum(WEIGHTS[x] for x in ["Food, Bvgs & Tobacco","Goods Ex-Food","Housing","Education","Other Services"]),
        "Non-core children": WEIGHTS["Agricultural"]+WEIGHTS["Energy & Controlled Prices"],
    }
    latest_b=bidx.index.max(); half=int(bidx.loc[latest_b,"half"])
    lines=[
        "Mexico INPC dashboard — QA report",
        "Generated: "+datetime.now().isoformat(timespec="seconds"),
        "",
        f"Monthly latest: {midx.index.max().strftime('%Y-%m')}",
        f"Biweekly latest: {('1H' if half==1 else '2H')} {latest_b.strftime('%Y-%m')}",
        f"Monthly SHA256: {file_sha256(monthly_path)}",
        f"Biweekly SHA256: {file_sha256(biweekly_path)}",
        "",
        "Structural QA: "+("PASS" if not issues else "WARN"),
    ]
    lines += ["- "+x for x in issues] if issues else ["- 16 components present; latest values complete; dates unique and ordered; index levels positive."]
    lines += ["", "Weight hierarchy checks:"]
    lines += [f"- {k}: {v:.6f}" for k,v in weight_checks.items()]
    lines += [
        "",
        "Contribution methodology:",
        "- For the latest releases available in the JSON cache, realized contribution uses INEGI's official valor_incidencia field.",
        "- Projected BTG contribution is the configured 2024 basket weight multiplied by the desk period forecast.",
        "- Surprise contribution is official realized incidence minus projected BTG contribution, converted to basis points.",
        "- If official incidence is unavailable for an older cached view, the HTML explicitly falls back to weight x realized period change as an analytical approximation.",
        "- YoY breakdown charts remain normalized analytical decompositions for readability and should not be described as official INEGI incidence.",
        "",
        "Seasonal adjustment:",
        "- STL on log index levels; period 12 for monthly / half-specific series and 24 for the full biweekly sequence; robust=True.",
        "",
        "Projection QA:",
    ]
    lines += ["- "+x for x in projection_warnings] if projection_warnings else ["- No projection-file warnings."]
    out=root/"logs"/"latest_qa_report.txt"
    out.parent.mkdir(parents=True,exist_ok=True)
    out.write_text("\n".join(lines),encoding="utf-8")
    return out


def main():
    ap=argparse.ArgumentParser()
    ap.add_argument("--offline",action="store_true",help="Skip live INEGI JSON query and rebuild from local historical seed/cache")
    ap.add_argument("--no-fallback",action="store_true",help="Fail instead of using the last valid local cache if the live INEGI JSON query fails")
    args=ap.parse_args()

    root=Path(__file__).resolve().parent
    monthly=root/"data"/"ca55_2018a.xlsx"
    biweekly=root/"data"/"ca56_2018a.xlsx"
    template=root/"dashboard_template.html"
    output=root/"output"/"mexico_inpc_dashboard.html"
    projections=root/"btg_projections.xlsx"

    if not args.offline:
        try:
            mlast,blast=update_release_cache(root)
            log(f"Live monthly release validated: {mlast}")
            log(f"Live biweekly release validated: {blast}")
        except Exception as exc:
            if args.no_fallback:
                raise
            log(f"WARNING: live INEGI JSON update failed: {exc}")
            log("Keeping the last valid JSON cache and rebuilding from local history.")

    validate_inegi_file(monthly,"monthly")
    validate_inegi_file(biweekly,"biweekly")
    midx,bidx,projection_warnings=rebuild_dashboard(template,output,monthly,biweekly,projections,root)
    qa=write_qa_report(root,monthly,biweekly,midx,bidx,projection_warnings)
    log(f"Dashboard generated: {output}")
    log(f"Latest monthly observation: {midx.index.max().strftime('%b %Y')}")
    latest_b=bidx.index.max(); h=int(bidx.loc[latest_b,'half'])
    log(f"Latest biweekly observation: {'1H' if h==1 else '2H'} {latest_b.strftime('%b %Y')}")
    log(f"QA report: {qa}")


if __name__=="__main__":
    main()
